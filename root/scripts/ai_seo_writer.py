# -*- coding: utf-8 -*-
import os
import sys
import openpyxl
import time
import re
from dotenv import load_dotenv
from openai import OpenAI

# ================= 1. 强行读取机密中心 =================
load_dotenv("/root/scripts/.env")

# 拿到 sub2api / new-api 的本地地址和万能钥匙
API_BASE_URL = os.getenv("API_BASE_URL")
API_KEY = os.getenv("API_KEY")

# 初始化统一的大模型客户端
client = OpenAI(
    api_key=API_KEY,
    base_url=API_BASE_URL,
    timeout=40.0
)

# ================= 2. 动态命令行传参 =================
PROCESS_LIMIT = 1
if len(sys.argv) > 1:
    arg = sys.argv[1].lower()
    if arg == 'all':
        PROCESS_LIMIT = 999999
    elif arg.isdigit():
        PROCESS_LIMIT = int(arg)
    else:
        print("⚠️ 参数错误！请使用: python3 ai_seo_writer.py [数字 或 all]")
        exit(1)

# ================= 3. 基础配置 =================
EXCEL_FILE = "/root/scripts/pantools_export.xlsx"
OUTPUT_FILE = "/root/scripts/pantools_export_seo.xlsx"
DB_FILE = os.getenv("DB_FILE", "/root/scripts/database.db")

# ✅ 你的高阶玩法：模型轮询列表（按优先级从高到低）
FALLBACK_MODELS = [
    "deepseek-v4-flash",
    "deepseek-v4-pro",
    "deepseek-v4-pro[1]"
]

# ================= 4. 数据清洗引擎 =================
def normalize_filename(name):
    if not name: return ""
    name = str(name).lower()
    noise_patterns = [
        r'\d{3,4}p', r'1080p', r'720p', r'4k', r'蓝光', r'高清', r'超清', 
        r'无删减', r'完整版', r'bd', r'web-dl', r'hd', r'x264', r'h265'
    ]
    for pattern in noise_patterns:
        name = re.sub(pattern, '', name)
    name = re.sub(r'[^a-z0-9\u4e00-\u9fa5]', '', name)
    return name.strip()

def clean_header_name(val):
    if not val: return ""
    return str(val).strip().replace(" ", "").replace("\n", "").replace("\r", "").lower()

# ================= 5. 核心调度与验证逻辑 =================
def validate_output(text):
    if not text: raise ValueError("返回内容为空")
    if len(text) < 300: raise ValueError(f"字数不足 ({len(text)} 字)")
    if "<p>" not in text and "<ul>" not in text: raise ValueError("未检测到 HTML 排版标签")
    
    english_chars = sum(1 for c in text if c.isalpha() and c.isascii())
    if english_chars > len(text) * 0.4:
         raise ValueError("检测到大段英文输出，拒绝接收！")
         
    forbidden_brands = ["夸克", "百度", "阿里云盘", "迅雷", "天翼", "115网盘", "移动云盘", "城通网盘"]
    for brand in forbidden_brands:
        if brand in text:
             raise ValueError(f"检测到违禁网盘品牌词 ({brand})，强行打回重写！")
             
    apology_words = ["作为一个人工智能", "抱歉", "无法提供", "我不能", "违反政策"]
    for word in apology_words:
        if word in text: raise ValueError(f"触发 AI 拒绝回复词汇 ({word})")
    return text

def generate_seo_desc(filename):
    # ✅ 核心改动：在喂给 AI 之前，先给文件名“洗个澡”，把敏感词剔除
    safe_filename = str(filename)
    forbidden_brands = ["夸克", "百度", "阿里云盘", "迅雷", "天翼", "115网盘", "移动云盘", "城通网盘"]
    for brand in forbidden_brands:
        safe_filename = safe_filename.replace(brand, "")
        
    prompt = f"""
你是一个资源分享站的编辑。请为资源《{safe_filename}》写一篇 500 字左右的介绍。

【写作要求】
1. 全程用简体中文。禁止说”作为一个人工智能”、”抱歉”这类话。
2. 不要提任何网盘品牌名，用”云盘””网盘”代替。
3. 语气像真人分享——直接、有料、不废话。不要营销口吻和emoji堆砌。
4. 写三件事：这份资料是什么、里面具体有什么、适合谁看能解决什么问题。
5. 避免雷同：同系列资源每篇的开场角度、亮点选取必须有明显差异，禁止套模板。
6. 写完读一遍，像机器人写的就重写。

【输出格式】直接输出 HTML：
<p>开场（80字，一句话说清这是什么，为什么值得下载）</p>
<p>核心内容（200字，具体说明里面有什么，不要空话）</p>
<ul><li>亮点1</li><li>亮点2</li><li>亮点3</li></ul>
<p>适合人群（100字，谁需要，看完能得到什么）</p>
<p>结尾（自然收尾，60字）</p>
"""
    # 在这个循环里依次尝试你的模型
    for model_name in FALLBACK_MODELS:
        try:
            print(f"   ⏳ 尝试调用模型 [{model_name}] ...")
            
            response = client.chat.completions.create(
                model=model_name,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.7,
                max_tokens=1500
            )
            
            validated_content = validate_output(response.choices[0].message.content)
            print(f"   ✅ [{model_name}] 成功接单并完成生成！")
            return validated_content
            
        except Exception as e:
            # 捕获 404, 503 等任何报错，打印日志后直接继续下一个循环
            print(f"   ⚠️ [{model_name}] 阵亡: {e}，正在切换下一个...")
            continue

    print("   🚨 所有备用模型均已阵亡，启用默认文本。")
    return "<p>优质极客资源，完美支持自动化环境部署，推荐进行网盘极速转存。</p>"

# ================= 6. 主流程 =================
print(f"🚀 启动引擎 | 本次任务限额: {'无限制 (ALL)' if PROCESS_LIMIT > 900000 else PROCESS_LIMIT} 条")
cache_dict = {}
if os.path.exists(OUTPUT_FILE):
    try:
        wb_cache = openpyxl.load_workbook(OUTPUT_FILE, data_only=True)
        sheet_cache = wb_cache.active
        header = {clean_header_name(c.value): i for i, c in enumerate(sheet_cache[1], 1) if c.value}
        if "文件名" in header and "seo简介" in header:
            for r in range(2, sheet_cache.max_row + 1):
                fname, fdesc = sheet_cache.cell(row=r, column=header["文件名"]).value, sheet_cache.cell(row=r, column=header["seo简介"]).value
                if fname and fdesc and len(str(fdesc).strip()) > 300 and "<p>" in str(fdesc):
                    english_ratio = sum(1 for c in str(fdesc) if c.isalpha() and c.isascii()) / len(str(fdesc))
                    if english_ratio < 0.4: 
                        cache_dict[normalize_filename(fname)] = str(fdesc).strip()
    except Exception:
        pass

wb = openpyxl.load_workbook(EXCEL_FILE)
sheet = wb.active
header = {clean_header_name(c.value): i for i, c in enumerate(sheet[1], 1) if c.value}

if "seo简介" not in header:
    seo_col = sheet.max_column + 1
    sheet.cell(row=1, column=seo_col, value="SEO简介")
else:
    seo_col = header["seo简介"]

file_col = header.get("文件名")
# ================= 连接数据库 =================
import sqlite3
db = sqlite3.connect(DB_FILE)
db.execute("PRAGMA journal_mode=WAL")

cache_hits, api_calls, processed_in_this_run = 0, 0, 0

for row in range(2, sheet.max_row + 1):
    if processed_in_this_run >= PROCESS_LIMIT:
        print(f"\n🛑 达到任务限额 ({PROCESS_LIMIT} 条)，程序平稳退出。")
        break

    filename = sheet.cell(row=row, column=file_col).value
    if not filename: continue

    filename_str = str(filename).strip()

    # Check if database already has good content
    db_row = db.execute("SELECT seo_desc FROM resources WHERE title=?", (filename_str,)).fetchone()
    if db_row and db_row[0] and len(str(db_row[0])) > 200 and "<p>" in str(db_row[0]):
        cache_hits += 1
        continue

    print(f"\n🤖 正在处理: {filename_str}")
    desc = generate_seo_desc(filename_str)

    # Write to database immediately
    db.execute("UPDATE resources SET seo_desc=? WHERE title=?", (desc, filename_str))
    db.commit()

    # Also save to Excel cache for backup
    sheet.cell(row=row, column=seo_col, value=desc)

    api_calls += 1
    processed_in_this_run += 1

    if api_calls % 5 == 0:
        wb.save(OUTPUT_FILE)
        db.commit()
    time.sleep(1)

wb.save(OUTPUT_FILE)
db.commit()
db.close()
print(f"========================================\n✅ 任务结束 | 🟢 DB缓存跳过: {cache_hits} | 🚀 AI生成: {api_calls}\n========================================")